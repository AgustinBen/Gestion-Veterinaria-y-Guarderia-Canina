"""Tests del lector de solo lectura del Excel histórico (T-15, Pasada A)."""

from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.modules.migracion.lector import (
    HOJAS_ESPERADAS,
    ErrorEstructuraExcel,
    ExcelNoEncontrado,
    leer_libro,
)
from app.modules.migracion.perfilado import (
    ErrorSinDatos,
    clasificar_fecha,
    construir_perfilado,
    convertir_edad_anios,
    convertir_peso_kg,
    es_agrupable_telefono,
    menciona_deuda,
    normalizar_telefono_temporal,
    perfilar,
    resumen_consola,
)
from scripts.perfilar_excel import main as main_perfilar


def _fecha_edicion(ruta: Path) -> datetime | None:
    libro = load_workbook(ruta, read_only=True)
    try:
        return libro.properties.modified
    finally:
        libro.close()


def test_lector_lee_las_tres_hojas_con_sus_encabezados(excel_ficticio: Path) -> None:
    libro = leer_libro(excel_ficticio)

    assert set(libro.hojas) == set(HOJAS_ESPERADAS)
    for nombre, columnas in HOJAS_ESPERADAS.items():
        hoja = libro.hojas[nombre]
        assert hoja.nombre == nombre
        assert hoja.encabezados == columnas


def test_lector_lee_las_filas_crudas_incluida_la_vacia(excel_ficticio: Path) -> None:
    libro = leer_libro(excel_ficticio)

    assert len(libro.hojas["Clientes_y_Mascotas"].filas) == 9
    assert len(libro.hojas["Atenciones_Clinicas"].filas) == 6
    assert len(libro.hojas["Guarderia_y_Estadias"].filas) == 5

    fila_vacia = libro.hojas["Clientes_y_Mascotas"].filas[-1]
    assert fila_vacia.numero == 10
    assert all(valor == "" for valor in fila_vacia.valores.values())


def test_lector_conserva_el_texto_original(excel_ficticio: Path) -> None:
    libro = leer_libro(excel_ficticio)
    fila = libro.hojas["Clientes_y_Mascotas"].filas[0]

    assert fila.numero == 2
    assert fila.valores["Telefono_Contacto"] == "+54 9 11 4455-8899"
    assert fila.valores["Especie_Raza"] == "Perro - Labrador"
    assert fila.valores["Peso_Kg"] == "28,5 kg"
    assert set(fila.valores) == set(HOJAS_ESPERADAS["Clientes_y_Mascotas"])


def test_lector_expone_la_fecha_de_ultima_edicion(excel_ficticio: Path) -> None:
    esperada = _fecha_edicion(excel_ficticio)
    libro = leer_libro(excel_ficticio)

    assert isinstance(esperada, datetime)
    assert libro.fecha_ultima_edicion == esperada


def test_lector_falla_si_falta_una_hoja(excel_ficticio: Path, tmp_path: Path) -> None:
    roto = tmp_path / "sin_guarderia.xlsx"
    libro = load_workbook(excel_ficticio)
    del libro["Guarderia_y_Estadias"]
    libro.save(roto)
    libro.close()

    with pytest.raises(ErrorEstructuraExcel, match="Guarderia_y_Estadias"):
        leer_libro(roto)


def test_lector_falla_si_renombran_una_columna(excel_ficticio: Path, tmp_path: Path) -> None:
    roto = tmp_path / "columna_renombrada.xlsx"
    libro = load_workbook(excel_ficticio)
    libro["Clientes_y_Mascotas"].cell(row=1, column=3).value = "Telefono"
    libro.save(roto)
    libro.close()

    with pytest.raises(ErrorEstructuraExcel) as excinfo:
        leer_libro(roto)

    mensaje = str(excinfo.value)
    assert "Telefono_Contacto" in mensaje
    assert "Clientes_y_Mascotas" in mensaje


def test_lector_falla_con_archivo_inexistente(tmp_path: Path) -> None:
    with pytest.raises(ExcelNoEncontrado, match="No se encontró"):
        leer_libro(tmp_path / "no_existe.xlsx")


def test_lector_no_modifica_el_archivo_de_origen(excel_ficticio: Path) -> None:
    antes = hashlib.sha256(excel_ficticio.read_bytes()).hexdigest()
    leer_libro(excel_ficticio)
    despues = hashlib.sha256(excel_ficticio.read_bytes()).hexdigest()

    assert antes == despues


def _seccion(datos: dict[str, object], nombre: str) -> dict[str, object]:
    valor = datos[nombre]
    assert isinstance(valor, dict)
    return valor


def _filas_csv(ruta: Path) -> list[dict[str, str]]:
    with ruta.open(encoding="utf-8-sig", newline="") as archivo:
        return list(csv.DictReader(archivo))


@pytest.fixture
def datos(excel_ficticio: Path) -> dict[str, object]:
    return construir_perfilado(leer_libro(excel_ficticio)).datos


@pytest.fixture
def salida(tmp_path: Path) -> Path:
    return tmp_path / "reportes"


@pytest.fixture
def reporte(excel_ficticio: Path, salida: Path):
    return perfilar(excel_ficticio, salida)


def test_conteo_filas_por_hoja(datos: dict[str, object]) -> None:
    conteo = _seccion(datos, "conteo_filas_por_hoja")

    assert conteo["Clientes_y_Mascotas"] == {
        "filas_totales": 10,
        "filas_hoja": 9,
        "filas_con_datos": 8,
    }
    assert conteo["Atenciones_Clinicas"] == {
        "filas_totales": 7,
        "filas_hoja": 6,
        "filas_con_datos": 6,
    }
    assert conteo["Guarderia_y_Estadias"] == {
        "filas_totales": 6,
        "filas_hoja": 5,
        "filas_con_datos": 5,
    }


def test_fechas_por_formato_y_por_columna(datos: dict[str, object]) -> None:
    fechas = _seccion(datos, "fechas")

    assert fechas["total_celdas"] == 16
    assert fechas["por_formato"] == {
        "dd/mm/aaaa": 4,
        "aaaa-mm-dd": 4,
        "d de Mes aaaa": 3,
        "dd/mm/aa": 3,
    }
    assert fechas["invalidas"] == 1
    assert fechas["no_reconocidas"] == 1
    assert fechas["vacias"] == 0

    columnas = _seccion(fechas, "por_columna")
    assert columnas["Atenciones_Clinicas.Fecha_Visita"] == {
        "total_celdas": 6,
        "vacias": 0,
        "por_formato": {
            "dd/mm/aaaa": 1,
            "aaaa-mm-dd": 1,
            "d de Mes aaaa": 1,
            "dd/mm/aa": 1,
        },
        "invalidas": 1,
        "no_reconocidas": 1,
    }
    assert columnas["Guarderia_y_Estadias.Fecha_Ingreso"] == {
        "total_celdas": 5,
        "vacias": 0,
        "por_formato": {
            "dd/mm/aaaa": 1,
            "aaaa-mm-dd": 2,
            "d de Mes aaaa": 1,
            "dd/mm/aa": 1,
        },
        "invalidas": 0,
        "no_reconocidas": 0,
    }
    assert columnas["Guarderia_y_Estadias.Fecha_Egreso"] == {
        "total_celdas": 5,
        "vacias": 0,
        "por_formato": {
            "dd/mm/aaaa": 2,
            "aaaa-mm-dd": 1,
            "d de Mes aaaa": 1,
            "dd/mm/aa": 1,
        },
        "invalidas": 0,
        "no_reconocidas": 0,
    }


def test_clasificar_fecha_acepta_las_dos_variantes_textuales() -> None:
    assert clasificar_fecha("12/03/2024").fecha == date(2024, 3, 12)
    assert clasificar_fecha("2024-05-04").fecha == date(2024, 5, 4)
    assert clasificar_fecha("7 de junio 2024").fecha == date(2024, 6, 7)
    assert clasificar_fecha("10 de octubre de 2024").fecha == date(2024, 10, 10)
    assert clasificar_fecha("05/06/24").fecha == date(2024, 6, 5)
    assert clasificar_fecha("31/02/2024").motivo == "fecha_invalida"
    assert clasificar_fecha("ayer").motivo == "fecha_no_reconocida"


def test_telefonos_conteos(datos: dict[str, object]) -> None:
    telefonos = _seccion(datos, "telefonos")

    assert telefonos["total_celdas"] == 8
    assert telefonos["vacias"] == 1
    assert telefonos["con_letras"] == 1
    assert telefonos["con_varios_numeros"] == 1
    assert telefonos["formatos_distintos"] == 7
    assert telefonos["agrupables"] == 5
    assert telefonos["con_digitos"] == 6
    assert telefonos["sin_digitos"] == 1
    assert len(telefonos["variantes"]) == 7
    assert all(variante["cantidad"] == 1 for variante in telefonos["variantes"])


def test_normalizacion_de_telefono_y_agrupabilidad() -> None:
    assert normalizar_telefono_temporal("+54 9 11 4455-8899") == "1144558899"
    assert normalizar_telefono_temporal("11 4455-8899") == "1144558899"
    assert normalizar_telefono_temporal("011 4890-1122") == "1148901122"
    assert normalizar_telefono_temporal("15-6677-2233") == "66772233"
    assert normalizar_telefono_temporal("341-555-0000") == "3415550000"
    assert normalizar_telefono_temporal("9114455889") == "9114455889"
    assert normalizar_telefono_temporal("541144558899") == "1144558899"

    assert es_agrupable_telefono("11 4455-8899")
    assert not es_agrupable_telefono("Sin telefono (dejo mail)")
    assert not es_agrupable_telefono("11 5555-1234 / 11 6666-5678")
    assert not es_agrupable_telefono("")


def test_edad_conteos_y_conversion(datos: dict[str, object]) -> None:
    edad = _seccion(datos, "edad")

    assert edad["total_celdas"] == 8
    assert edad["vacias"] == 0
    assert edad["reconocidas"] == 7
    assert edad["no_reconocidas"] == 1
    assert edad["por_anio"] == {
        "0": 1,
        "2": 1,
        "3": 1,
        "4": 2,
        "8": 1,
        "10": 1,
    }

    assert convertir_edad_anios("4 años") == 4
    assert convertir_edad_anios("4") == 4
    assert convertir_edad_anios("3 años aprox") == 3
    assert convertir_edad_anios("2 años y medio") == 2
    assert convertir_edad_anios("6 meses") == 0
    assert convertir_edad_anios("no sé") is None


def test_peso_conteos_y_conversion(datos: dict[str, object]) -> None:
    peso = _seccion(datos, "peso")

    assert peso["total_celdas"] == 8
    assert peso["vacias"] == 1
    assert peso["con_unidad"] == 5
    assert peso["sin_unidad"] == 2
    assert peso["separador_coma"] == 3
    assert peso["separador_punto"] == 2
    assert peso["validos"] == 7
    assert peso["no_reconocidos"] == 0

    parseado = convertir_peso_kg("28,5 kg")
    assert parseado is not None
    assert parseado.valor == Decimal("28.5")
    assert parseado.con_unidad
    assert parseado.separador == ","

    parseado = convertir_peso_kg("3.5")
    assert parseado is not None
    assert parseado.valor == Decimal("3.5")
    assert not parseado.con_unidad
    assert parseado.separador == "."


def test_especie_raza_conteos(datos: dict[str, object]) -> None:
    especie = _seccion(datos, "especie_raza")

    assert especie["total_celdas"] == 8
    assert especie["vacias"] == 0
    assert especie["valores_distintos"] == 7
    assert especie["normalizados_distintos"] == 6
    assert especie["grupos_variantes"] == 1
    assert especie["variantes"] == [
        {
            "normalizado": "perro - labrador",
            "originales": ["Perro - Labrador", "perro - labrador"],
            "cantidad": 2,
        }
    ]


def test_duplicados(datos: dict[str, object]) -> None:
    duplicados = _seccion(datos, "duplicados")

    assert duplicados["por_telefono"]["grupos"] == 1
    assert duplicados["por_telefono"]["filas"] == 2
    assert duplicados["por_duenio"]["grupos"] == 2
    assert duplicados["por_duenio"]["filas"] == 4
    assert duplicados["por_mascota"]["grupos"] == 1
    assert duplicados["por_mascota"]["filas"] == 2


def test_mascotas_sin_duenio_por_par(datos: dict[str, object], reporte, salida: Path) -> None:
    huerfanos = _seccion(datos, "mascotas_sin_duenio")

    assert huerfanos["pares_unicos"] == 1
    assert huerfanos["filas"] == 2
    assert huerfanos["detalle"] == [
        {"hoja": "Atenciones_Clinicas", "fila": 6},
        {"hoja": "Guarderia_y_Estadias", "fila": 6},
    ]

    filas = [
        fila
        for fila in _filas_csv(salida / "perfilado_filas.csv")
        if fila["motivo"] == "mascota_sin_duenio"
    ]
    assert [fila["valor_original"] for fila in filas] == ["Firulais", "Firulais"]
    assert "Chicho" not in {fila["valor_original"] for fila in filas}
    assert "Titan" not in {fila["valor_original"] for fila in filas}
    assert "Rocco Jr." not in {fila["valor_original"] for fila in filas}


def test_personas_sin_apellido(datos: dict[str, object], reporte, salida: Path) -> None:
    sin_apellido = _seccion(datos, "sin_apellido")

    assert sin_apellido["personas_unicas"] == 3
    assert sin_apellido["duenios_tutores_unicos"] == 2
    assert sin_apellido["veterinarios_unicos"] == 1
    assert sin_apellido["filas"] == 5

    valores = [
        fila["valor_original"]
        for fila in _filas_csv(salida / "perfilado_filas.csv")
        if fila["motivo"] == "persona_sin_apellido"
    ]
    assert set(valores) == {"Lucia", "PEDRO", "Dra. Laura"}


def test_fila_vacia_no_contamina_columnas(datos: dict[str, object]) -> None:
    vacias = _seccion(datos, "filas_totalmente_vacias")
    assert vacias["cantidad"] == 1
    assert vacias["detalle"] == [{"hoja": "Clientes_y_Mascotas", "fila": 10}]

    assert _seccion(datos, "telefonos")["total_celdas"] == 8
    assert _seccion(datos, "edad")["total_celdas"] == 8
    assert _seccion(datos, "peso")["total_celdas"] == 8
    assert _seccion(datos, "especie_raza")["total_celdas"] == 8


def test_deuda_en_texto(datos: dict[str, object], reporte, salida: Path) -> None:
    deuda = _seccion(datos, "deuda_en_el_texto")

    assert deuda["filas"] == 3
    assert deuda["detalle"] == [
        {"hoja": "Atenciones_Clinicas", "fila": 6},
        {"hoja": "Guarderia_y_Estadias", "fila": 3},
        {"hoja": "Guarderia_y_Estadias", "fila": 4},
    ]

    assert menciona_deuda("Debe $3000")
    assert menciona_deuda("Fiado")
    assert menciona_deuda("Seña pendiente")
    assert not menciona_deuda("Saldo: $0")
    assert not menciona_deuda("Pagado")

    valores = [
        fila["valor_original"]
        for fila in _filas_csv(salida / "perfilado_filas.csv")
        if fila["motivo"] == "deuda_en_texto"
    ]
    assert valores == ["Debe $5000", "Debe $3000", "Fiado"]


def test_dni_solo_en_json(datos: dict[str, object], reporte, salida: Path) -> None:
    dni = _seccion(datos, "dni")
    assert dni["columna_presente"] is False
    assert "hallazgo" in dni

    motivos = {fila["motivo"] for fila in _filas_csv(salida / "perfilado_filas.csv")}
    assert "dni_ausente" not in motivos


def test_salidas_fecha_y_csv(reporte, salida: Path, excel_ficticio: Path) -> None:
    modificado = _fecha_edicion(excel_ficticio)
    assert isinstance(modificado, datetime)
    assert reporte.datos["fecha_ultima_edicion"] == modificado.isoformat()
    assert reporte.datos["fecha_corte_propuesta"] == modificado.date().isoformat()

    json_destino = salida / "perfilado.json"
    csv_destino = salida / "perfilado_filas.csv"
    assert json_destino.is_file()
    assert csv_destino.is_file()
    assert json.loads(json_destino.read_text(encoding="utf-8"))["dni"]["columna_presente"] is False
    assert csv_destino.read_bytes()[:3] == b"\xef\xbb\xbf"

    with csv_destino.open(encoding="utf-8-sig", newline="") as archivo:
        lector = csv.DictReader(archivo)
        assert lector.fieldnames == ["hoja", "fila", "columna", "motivo", "valor_original"]
        filas = list(lector)

    assert len(filas) == 25
    assert Counter(fila["motivo"] for fila in filas) == Counter(
        {
            "fila_vacia": 1,
            "telefono_vacio": 1,
            "telefono_con_letras": 1,
            "telefono_con_varios_numeros": 1,
            "fecha_invalida": 1,
            "fecha_no_reconocida": 1,
            "edad_no_reconocida": 1,
            "persona_sin_apellido": 5,
            "duplicado_telefono": 2,
            "duplicado_duenio": 4,
            "duplicado_mascota": 2,
            "mascota_sin_duenio": 2,
            "deuda_en_texto": 3,
        }
    )


def test_determinismo_y_hash(excel_ficticio: Path, tmp_path: Path) -> None:
    antes = hashlib.sha256(excel_ficticio.read_bytes()).hexdigest()
    salida_1 = tmp_path / "salida_1"
    salida_2 = tmp_path / "salida_2"

    perfilar(excel_ficticio, salida_1)
    perfilar(excel_ficticio, salida_2)

    assert (salida_1 / "perfilado.json").read_bytes() == (salida_2 / "perfilado.json").read_bytes()
    assert (salida_1 / "perfilado_filas.csv").read_bytes() == (
        salida_2 / "perfilado_filas.csv"
    ).read_bytes()
    assert antes == hashlib.sha256(excel_ficticio.read_bytes()).hexdigest()


def test_resumen_consola_sin_datos_personales(reporte) -> None:
    resumen = resumen_consola(reporte)

    assert "Perfilado completado." in resumen
    assert "Hallazgos: 25" in resumen
    assert "VALERIA" not in resumen
    assert "1144558899" not in resumen
    assert "Firulais" not in resumen


def test_script_genera_reporte(excel_ficticio: Path, tmp_path: Path, capsys) -> None:
    salida = tmp_path / "reportes"
    codigo = main_perfilar(
        {"excel_origen": str(excel_ficticio), "migracion_salida": str(salida)}
    )

    assert codigo == 0
    assert (salida / "perfilado.json").is_file()
    assert (salida / "perfilado_filas.csv").is_file()
    capturado = capsys.readouterr()
    assert "Perfilado completado." in capturado.out
    assert "VALERIA" not in capturado.out


def test_script_archivo_inexistente(tmp_path: Path, capsys) -> None:
    codigo = main_perfilar(
        {
            "excel_origen": str(tmp_path / "no_existe.xlsx"),
            "migracion_salida": str(tmp_path / "reportes"),
        }
    )

    assert codigo == 3
    assert "Error de entrada" in capsys.readouterr().err


def test_script_hoja_faltante(excel_ficticio: Path, tmp_path: Path, capsys) -> None:
    roto = tmp_path / "sin_hoja.xlsx"
    libro = load_workbook(excel_ficticio)
    del libro["Atenciones_Clinicas"]
    libro.save(roto)
    libro.close()

    codigo = main_perfilar(
        {"excel_origen": str(roto), "migracion_salida": str(tmp_path / "reportes")}
    )

    assert codigo == 2
    assert "Atenciones_Clinicas" in capsys.readouterr().err


def test_script_columna_renombrada(excel_ficticio: Path, tmp_path: Path, capsys) -> None:
    roto = tmp_path / "columna_renombrada.xlsx"
    libro = load_workbook(excel_ficticio)
    libro["Clientes_y_Mascotas"].cell(row=1, column=3).value = "Telefono"
    libro.save(roto)
    libro.close()

    codigo = main_perfilar(
        {"excel_origen": str(roto), "migracion_salida": str(tmp_path / "reportes")}
    )

    assert codigo == 2
    assert "Telefono_Contacto" in capsys.readouterr().err


def test_libro_sin_datos(excel_ficticio: Path, tmp_path: Path, capsys) -> None:
    sin_datos = tmp_path / "sin_datos.xlsx"
    libro = load_workbook(excel_ficticio)
    for hoja in libro.worksheets:
        hoja.delete_rows(2, hoja.max_row - 1)
    libro.save(sin_datos)
    libro.close()

    with pytest.raises(ErrorSinDatos):
        construir_perfilado(leer_libro(sin_datos))

    salida = tmp_path / "reportes"
    codigo = main_perfilar(
        {"excel_origen": str(sin_datos), "migracion_salida": str(salida)}
    )

    assert codigo == 4
    assert not (salida / "perfilado.json").exists()
    assert "ninguna fila con datos" in capsys.readouterr().err
