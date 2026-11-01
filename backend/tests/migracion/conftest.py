"""Fixtures del módulo de migración: Excel ficticio con defectos conocidos."""

from __future__ import annotations

import csv
from collections.abc import Callable, Generator, Sequence
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.engine import Engine
from sqlalchemy.orm import Session

from app.core import parametros
from app.modules.migracion import carga
from app.modules.migracion.carga import ReporteCarga
from app.modules.migracion.models import MigracionRegistro
from app.modules.migracion.normalizacion import NormalizadorTelefono
from app.modules.migracion.seed import (
    CLAVE_FECHA_CORTE,
    CLAVE_PESO_MAX_KG,
    sembrar_parametros_migracion,
)
from app.modules.padron import models as padron_models
from app.modules.padron import service as padron_service
from app.modules.padron.seed import (
    sembrar_especies_tamanios_tipo_contacto,
    sembrar_parametros_padron,
)

COLUMNAS_CLIENTES = (
    "ID_Registro",
    "Nombre_Duenio",
    "Telefono_Contacto",
    "Nombre_Mascota",
    "Especie_Raza",
    "Edad_Aprox",
    "Peso_Kg",
    "Direccion_Localidad",
)

COLUMNAS_ATENCIONES = (
    "Nro_Ficha",
    "Fecha_Visita",
    "Paciente_Mascota",
    "Duenio_Referencia",
    "Tipo_Servicio",
    "Detalle_Clinico_y_Cobro",
    "Veterinario_Tratante",
)

COLUMNAS_GUARDERIA = (
    "Nro_Reserva",
    "Mascota",
    "Tutor",
    "Fecha_Ingreso",
    "Fecha_Egreso",
    "Modalidad",
    "Tarifa_Total",
    "Estado_Pago",
    "Comportamiento_Cuidados",
)

FILAS_CLIENTES = (
    (
        "REG-001",
        "VALERIA GOMEZ",
        "+54 9 11 4455-8899",
        "Firulais",
        "Perro - Labrador",
        "4 años",
        "28,5 kg",
        "Av. Siempre Viva 742, CABA",
    ),
    (
        "REG-002",
        "valeria gomez",
        "11 4455-8899",
        "Firulais",
        "perro - labrador",
        "4",
        "28.5kg",
        "Av. Siempre Viva 742, CABA",
    ),
    (
        "REG-003",
        "MARCOS PEREZ",
        "15-6677-2233",
        "Chicho",
        "Gato - Siames",
        "3 años aprox",
        "4 kg",
        "San Martin 100, CABA",
    ),
    (
        "REG-004",
        "Lucia",
        "Sin telefono (dejo mail)",
        "Titan",
        "canino",
        "6 meses",
        "",
        "Rivadavia 200, CABA",
    ),
    (
        "REG-005",
        "ANA LOPEZ",
        "",
        "Rocco Jr.",
        "Mestizo mediano",
        "2 años y medio",
        "12,25",
        "Belgrano 300, CABA",
    ),
    (
        "REG-006",
        "ANA LOPEZ",
        "011 4890-1122",
        "Rocco",
        "Mestizo mediano",
        "no sé",
        "12 kg",
        "Belgrano 300, CABA",
    ),
    (
        "REG-007",
        "PEDRO",
        "11 5555-1234 / 11 6666-5678",
        "Luna",
        "Gato Siames",
        "10",
        "3.5",
        "Mitre 400, Rosario",
    ),
    (
        "REG-008",
        "CARLA DIAZ",
        "341-555-0000",
        "Nala",
        "Perro - Caniche",
        "8 años",
        "6,4 kg",
        "Cordoba 500, Rosario",
    ),
    ("", "", "", "", "", "", "", ""),
)

FILAS_ATENCIONES = (
    (
        "A-001",
        "12/03/2024",
        "Firulais",
        "VALERIA GOMEZ",
        "Consulta",
        "Control sano",
        "Dr. Juan Perez",
    ),
    (
        "A-002",
        "2024-05-04",
        "Chicho",
        "MARCOS PEREZ",
        "Vacunación",
        "Antirrábica",
        "Dra. Laura",
    ),
    (
        "A-003",
        "7 de junio 2024",
        "Firulais",
        "VALERIA GOMEZ",
        "Consulta",
        "Control",
        "Dr. Juan Perez",
    ),
    (
        "A-004",
        "05/06/24",
        "Titan",
        "Lucia",
        "Cirugía",
        "Castración",
        "Dra. Laura",
    ),
    (
        "A-005",
        "ayer",
        "Firulais",
        "JULIA SOSA",
        "Consulta",
        "Debe $5000",
        "Dra. Laura",
    ),
    (
        "A-006",
        "31/02/2024",
        "Luna",
        "PEDRO",
        "Consulta",
        "Control",
        "Dr. Juan Perez",
    ),
)

FILAS_GUARDERIA = (
    (
        "G-001",
        "Rocco Jr.",
        "ANA LOPEZ",
        "2024-07-01",
        "03/07/2024",
        "Día",
        "15000",
        "Pagado",
        "Tranquilo",
    ),
    (
        "G-002",
        "Rocco",
        "ANA LOPEZ",
        "15/07/24",
        "17/07/24",
        "Noche",
        "15000",
        "Debe $3000",
        "Juguetón",
    ),
    (
        "G-003",
        "Chicho",
        "MARCOS PEREZ",
        "2024-08-01",
        "2024-08-03",
        "Día",
        "12000",
        "Fiado",
        "Tranquilo",
    ),
    (
        "G-004",
        "Titan",
        "Lucia",
        "01/09/2024",
        "02/09/2024",
        "Noche",
        "12000",
        "Pagado",
        "",
    ),
    (
        "G-005",
        "Firulais",
        "JULIA SOSA",
        "10 de octubre de 2024",
        "12 de octubre de 2024",
        "Día",
        "10000",
        "Pagado",
        "Tranquilo",
    ),
)


def _agregar_hoja(libro: Workbook, nombre: str, columnas: tuple[str, ...], filas: tuple[tuple[str, ...], ...]) -> None:
    hoja = libro.create_sheet(title=nombre)
    hoja.append(columnas)
    for fila in filas:
        hoja.append(fila)


def _crear_excel(ruta: Path) -> None:
    libro = Workbook()
    libro.remove(libro.active)
    _agregar_hoja(libro, "Clientes_y_Mascotas", COLUMNAS_CLIENTES, FILAS_CLIENTES)
    _agregar_hoja(libro, "Atenciones_Clinicas", COLUMNAS_ATENCIONES, FILAS_ATENCIONES)
    _agregar_hoja(libro, "Guarderia_y_Estadias", COLUMNAS_GUARDERIA, FILAS_GUARDERIA)
    libro.save(ruta)
    libro.close()


@pytest.fixture
def excel_ficticio(tmp_path: Path) -> Path:
    ruta = tmp_path / "excel_ficticio.xlsx"
    _crear_excel(ruta)
    return ruta


@pytest.fixture
def normalizador_telefono(
    db_session: Session,
) -> Generator[NormalizadorTelefono, None, None]:
    """Normalizador oficial de padrón, con su configuración sembrada."""
    parametros._limpiar_cache()
    sembrar_parametros_padron(db_session)
    yield lambda texto: padron_service.normalizar_telefono(db_session, texto)
    parametros._limpiar_cache()


# --- Apoyo de los tests de carga (T-16) ---------------------------------------

HOJA_CLIENTES = "Clientes_y_Mascotas"
HOJAS = (
    (HOJA_CLIENTES, COLUMNAS_CLIENTES),
    ("Atenciones_Clinicas", COLUMNAS_ATENCIONES),
    ("Guarderia_y_Estadias", COLUMNAS_GUARDERIA),
)
FECHA_CORTE = "2026-11-01"
PESO_MAX_KG = Decimal("60")


def crear_excel(
    ruta: Path,
    filas: Sequence[Sequence[str]],
    *,
    columnas: Sequence[str] = COLUMNAS_CLIENTES,
    hoja: str = HOJA_CLIENTES,
) -> Path:
    """Arma un .xlsx con las filas pedidas; las otras hojas van con encabezados."""
    libro = Workbook()
    libro.remove(libro.active)
    for nombre, encabezados in HOJAS:
        destino = libro.create_sheet(title=nombre)
        if nombre == hoja:
            destino.append(list(columnas))
            for fila in filas:
                destino.append(list(fila))
        else:
            destino.append(list(encabezados))
    libro.save(ruta)
    libro.close()
    return ruta


@pytest.fixture
def fila_cliente() -> Callable[..., tuple[str, ...]]:
    """Fila de Clientes_y_Mascotas con valores válidos que cada test ajusta."""

    def _fila(
        id_registro: str = "REG-001",
        nombre: str = "VALERIA GOMEZ",
        telefono: str = "11 4455-8899",
        mascota: str = "Firulais",
        especie: str = "Perro - Mestizo",
        edad: str = "3 años",
        peso: str = "10 kg",
        direccion: str = "Calle 1, CABA",
    ) -> tuple[str, ...]:
        return (id_registro, nombre, telefono, mascota, especie, edad, peso, direccion)

    return _fila


@pytest.fixture
def entorno_migracion(db_session: Session) -> Generator[None, None, None]:
    """Catálogos y parámetros con ``commit()``: la carga escribe por otra conexión."""
    parametros._limpiar_cache()
    sembrar_especies_tamanios_tipo_contacto(db_session)
    sembrar_parametros_migracion(db_session)
    db_session.commit()
    parametros._limpiar_cache()
    yield
    parametros._limpiar_cache()


@pytest.fixture
def fecha_corte(db_session: Session, entorno_migracion: None) -> None:
    """Fecha de corte y peso máximo de prueba, cargados en configuración."""
    parametros.set_param(
        db_session, CLAVE_FECHA_CORTE, FECHA_CORTE, "str", modulo="migracion"
    )
    parametros.set_param(
        db_session, CLAVE_PESO_MAX_KG, PESO_MAX_KG, "decimal", modulo="migracion"
    )
    parametros._limpiar_cache()


@pytest.fixture
def cargar(
    engine_sqlite: Engine,
    tmp_path: Path,
    entorno_migracion: None,
) -> Callable[..., ReporteCarga]:
    """Corre la carga contra la base en archivo del test."""

    def _cargar(
        excel: Path,
        *,
        simular: bool = False,
        salida: Path | None = None,
    ) -> ReporteCarga:
        return carga.cargar_padron(
            ruta_excel=excel,
            salida=tmp_path / "salida" if salida is None else salida,
            engine=engine_sqlite,
            simular=simular,
        )

    return _cargar


@pytest.fixture
def contar(engine_sqlite: Engine) -> Callable[[type], int]:
    """Cuenta filas en una sesión nueva, para ver lo que la carga commiteó."""

    def _contar(modelo: type) -> int:
        with Session(bind=engine_sqlite) as db:
            return db.scalar(select(func.count()).select_from(modelo)) or 0

    return _contar


@pytest.fixture
def registro(engine_sqlite: Engine) -> Callable[[int], dict[str, object] | None]:
    """Estado guardado de una fila del Excel, leído en sesión nueva."""

    def _registro(fila: int) -> dict[str, object] | None:
        with Session(bind=engine_sqlite) as db:
            encontrado = db.execute(
                select(MigracionRegistro).where(MigracionRegistro.fila == fila)
            ).scalar_one_or_none()
            if encontrado is None:
                return None
            return {
                "estado": encontrado.estado,
                "motivo": encontrado.motivo,
                "entidad": encontrado.entidad,
                "entidad_id": encontrado.entidad_id,
                "hash_fila": encontrado.hash_fila,
                "texto_original": encontrado.texto_original,
            }

    return _registro


@pytest.fixture
def personas(engine_sqlite: Engine) -> Callable[[], list[dict[str, object]]]:
    def _personas() -> list[dict[str, object]]:
        with Session(bind=engine_sqlite) as db:
            return [
                {"id": persona.id, "nombre": persona.nombre, "apellido": persona.apellido}
                for persona in db.execute(
                    select(padron_models.Persona).order_by(padron_models.Persona.id)
                ).scalars()
            ]

    return _personas


@pytest.fixture
def pacientes(engine_sqlite: Engine) -> Callable[[], list[dict[str, object]]]:
    def _pacientes() -> list[dict[str, object]]:
        with Session(bind=engine_sqlite) as db:
            return [
                {"id": paciente.id, "nombre": paciente.nombre, "peso_kg": paciente.peso_kg}
                for paciente in db.execute(
                    select(padron_models.Paciente).order_by(padron_models.Paciente.id)
                ).scalars()
            ]

    return _pacientes


@pytest.fixture
def leer_revision(tmp_path: Path) -> Callable[..., list[dict[str, str]]]:
    """Lee el CSV de revisión de la corrida (real o simulada)."""

    def _leer(
        salida: Path | None = None, *, simular: bool = False
    ) -> list[dict[str, str]]:
        directorio = tmp_path / "salida" if salida is None else salida
        nombre = "revision_simulada.csv" if simular else "revision.csv"
        with (directorio / nombre).open(encoding="utf-8-sig", newline="") as archivo:
            return list(csv.DictReader(archivo))

    return _leer

