"""Lectura de solo lectura del Excel histórico de migración.

No interpreta los datos: devuelve el texto de cada celda y valida la estructura
mínima (hojas y columnas) para que el perfilado nunca genere un reporte vacío
por un archivo equivocado.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

HOJAS_ESPERADAS: dict[str, tuple[str, ...]] = {
    "Clientes_y_Mascotas": (
        "ID_Registro",
        "Nombre_Duenio",
        "Telefono_Contacto",
        "Nombre_Mascota",
        "Especie_Raza",
        "Edad_Aprox",
        "Peso_Kg",
        "Direccion_Localidad",
    ),
    "Atenciones_Clinicas": (
        "Nro_Ficha",
        "Fecha_Visita",
        "Paciente_Mascota",
        "Duenio_Referencia",
        "Tipo_Servicio",
        "Detalle_Clinico_y_Cobro",
        "Veterinario_Tratante",
    ),
    "Guarderia_y_Estadias": (
        "Nro_Reserva",
        "Mascota",
        "Tutor",
        "Fecha_Ingreso",
        "Fecha_Egreso",
        "Modalidad",
        "Tarifa_Total",
        "Estado_Pago",
        "Comportamiento_Cuidados",
    ),
}


class ErrorLecturaExcel(Exception):
    """Error de lectura o de estructura del libro de migración."""


class ExcelNoEncontrado(ErrorLecturaExcel):
    """El archivo no existe o no se puede abrir."""


class ErrorEstructuraExcel(ErrorLecturaExcel):
    """Falta una hoja o una columna requerida."""


@dataclass(frozen=True)
class FilaCruda:
    numero: int
    valores: dict[str, str]


@dataclass(frozen=True)
class HojaCruda:
    nombre: str
    encabezados: tuple[str, ...]
    filas: tuple[FilaCruda, ...]


@dataclass(frozen=True)
class LibroCrudo:
    ruta: Path
    fecha_ultima_edicion: datetime | None
    hojas: dict[str, HojaCruda]


def _texto(valor: object) -> str:
    """Convierte una celda a texto sin interpretar su contenido."""
    if valor is None:
        return ""
    return str(valor)


def _validar_encabezados(
    hoja: str,
    encontrados: tuple[str, ...],
    esperados: tuple[str, ...],
) -> None:
    faltantes = [columna for columna in esperados if columna not in encontrados]
    if not faltantes:
        return

    detalle_faltantes = ", ".join(faltantes)
    detalle_encontrados = ", ".join(encontrados) or "(sin encabezados)"
    raise ErrorEstructuraExcel(
        f"La hoja '{hoja}' no tiene las columnas requeridas: {detalle_faltantes}. "
        f"Columnas encontradas: {detalle_encontrados}."
    )


def leer_libro(ruta: Path) -> LibroCrudo:
    """Lee las hojas requeridas con openpyxl en modo solo lectura.

    No modifica el archivo ni convierte los valores de celda a tipos de negocio.
    """
    ruta = Path(ruta)
    if not ruta.is_file():
        raise ExcelNoEncontrado(f"No se encontró el archivo Excel: {ruta}")

    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl lanza errores variados al abrir un archivo inválido
        raise ExcelNoEncontrado(f"No se pudo abrir el archivo Excel '{ruta}': {exc}") from exc

    try:
        modificado = libro.properties.modified
        fecha_ultima_edicion = modificado if isinstance(modificado, datetime) else None

        hojas: dict[str, HojaCruda] = {}
        for nombre, columnas in HOJAS_ESPERADAS.items():
            if nombre not in libro.sheetnames:
                encontradas = ", ".join(libro.sheetnames) or "(sin hojas)"
                raise ErrorEstructuraExcel(
                    f"Falta la hoja requerida '{nombre}'. Hojas encontradas: {encontradas}."
                )

            hoja = libro[nombre]
            filas_crudas = hoja.iter_rows(values_only=True)
            encabezados_crudos = next(filas_crudas, ())
            encabezados = tuple(_texto(valor).strip() for valor in encabezados_crudos)
            _validar_encabezados(nombre, encabezados, columnas)

            indice = {columna: posicion for posicion, columna in enumerate(encabezados)}
            filas: list[FilaCruda] = []
            for numero, valores in enumerate(filas_crudas, start=2):
                valores_fila = tuple(valores)
                por_columna = {
                    columna: _texto(
                        valores_fila[indice[columna]]
                        if indice[columna] < len(valores_fila)
                        else None
                    )
                    for columna in columnas
                }
                filas.append(FilaCruda(numero=numero, valores=por_columna))

            hojas[nombre] = HojaCruda(
                nombre=nombre,
                encabezados=encabezados,
                filas=tuple(filas),
            )

        return LibroCrudo(
            ruta=ruta,
            fecha_ultima_edicion=fecha_ultima_edicion,
            hojas=hojas,
        )
    finally:
        libro.close()
